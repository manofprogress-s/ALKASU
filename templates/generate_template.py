"""ALKASU içe aktarma şablonunu üretir: templates/urun_ice_aktarma_sablonu.xlsx ve .csv"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = "/home/claude/alkasu/templates/"
F = "Arial"
HEAD_REQ = PatternFill("solid", fgColor="1F4E79")
HEAD_OPT = PatternFill("solid", fgColor="5B7FA6")
INPUT = PatternFill("solid", fgColor="FFF9DB")
thin = Side(style="thin", color="C9CED6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# (anahtar, zorunlu, genişlik, açıklama, tür)  tür: text|int|money|pct|yn|list:<ad>
PRODUCTS = [
    ("urun_kodu", True, 14, "Benzersiz kod. Aynı kod tekrar yüklenirse ürün güncellenir.", "text"),
    ("urun_adi", True, 34, "Satış ekranında görünen ad. Ör: Erikli 19 L Damacana", "text"),
    ("marka", True, 16, "Yoksa otomatik oluşturulur.", "text"),
    ("kategori", True, 16, "Ör: Su, Damacana, Maden suyu, Meşrubat, Boş kap. Yoksa otomatik oluşturulur.", "text"),
    ("temel_birim", True, 12, "Stoğun tutulduğu en küçük birim. Ör: Adet, Şişe, Damacana", "list:birim"),
    ("barkod", False, 18, "Birden fazla barkod | ile ayrılır. Ör: 8690000000011|8690000000028", "text"),
    ("satis_fiyati", False, 12, "Temel birimin KDV dahil satış fiyatı (₺). Boş kap ürünlerinde boş bırakılır.", "money"),
    ("paket_adi", False, 11, "1. ek birimin adı (varsayılan: Paket). Kullanılmıyorsa boş.", "text"),
    ("paket_icerik", False, 11, "1 paket kaç temel birim? Tam sayı, 1'den büyük.", "int"),
    ("paket_satis_fiyati", False, 13, "KDV dahil paket fiyatı (₺). Boşsa paket satışta seçilemez.", "money"),
    ("koli_adi", False, 11, "2. ek birimin adı (varsayılan: Koli). Kullanılmıyorsa boş.", "text"),
    ("koli_icerik", False, 11, "1 koli kaç temel birim? Tam sayı, 1'den büyük.", "int"),
    ("koli_satis_fiyati", False, 13, "KDV dahil koli fiyatı (₺). Boşsa koli satışta seçilemez.", "money"),
    ("kdv_orani", True, 9, "0, 1, 10 veya 20", "list:kdv"),
    ("alis_fiyati", False, 12, "Temel birimin KDV dahil alış fiyatı (₺). Açılış maliyeti olur. Yalnızca yönetici görür.", "money"),
    ("acilis_stogu", False, 12, "Temel birim cinsinden eldeki miktar. Yalnızca ürünün hiç hareketi yoksa uygulanır.", "int"),
    ("kritik_stok", False, 11, "Temel birim. Bu seviyenin altına inince uyarı verilir.", "int"),
    ("depozitolu", False, 11, "E = dolu damacana gibi depozitolu ürün. Varsayılan H.", "list:eh"),
    ("depozito_tutari", False, 13, "Depozitolu üründe kap başına depozito (₺).", "money"),
    ("bos_kap_urun_kodu", False, 16, "Depozitolu ürünün bağlı olduğu boş kap ürününün kodu (aynı dosyada tanımlı olmalı).", "text"),
    ("bos_kap_mi", False, 10, "E = bu satır bir boş kap ürünü (ör. Erikli Boş Damacana). Satılmaz, yalnızca stoğu izlenir.", "list:eh"),
    ("aktif", False, 8, "E veya H. Varsayılan E.", "list:eh"),
]

CUSTOMERS = [
    ("musteri_kodu", True, 14, "Benzersiz kod. Aynı kod tekrar yüklenirse müşteri güncellenir.", "text"),
    ("ad_unvan", True, 30, "Müşteri adı veya firma unvanı", "text"),
    ("telefon", False, 15, "Ör: 05321234567", "text"),
    ("adres", False, 36, "Teslimat adresi / not", "text"),
    ("vergi_no", False, 14, "İsteğe bağlı", "text"),
    ("kredi_limiti", False, 12, "Veresiye limiti (₺). Boş veya 0 = veresiye kapalı.", "money"),
    ("limitsiz", False, 9, "E = limitsiz veresiye. Varsayılan H.", "list:eh"),
    ("acilis_bakiyesi", False, 14, "Geçişte müşterinin size olan veresiye borcu (₺). Pozitif = müşteri borçlu.", "money"),
]

CONTAINERS = [
    ("musteri_kodu", True, 14, "Müşteriler sayfasındaki kod", "text"),
    ("dolu_urun_kodu", True, 16, "Depozitolu ürünün kodu (ör. dolu damacana)", "text"),
    ("kap_adedi", True, 10, "Müşterinin elinde bulunan sizin kaplarınızın sayısı", "int"),
    ("odenen_depozito", False, 14, "Müşterinin bu kaplar için ödediği toplam depozito (₺). Boşsa 0.", "money"),
]

LISTS = {
    "birim": ["Adet", "Şişe", "Damacana", "Kutu", "Paket", "Bidon"],
    "kdv": [0, 1, 10, 20],
    "eh": ["E", "H"],
}

wb = Workbook()

# ---- Listeler (gizli)
ls = wb.active
ls.title = "Listeler"
for ci, (name, vals) in enumerate(LISTS.items(), start=1):
    ls.cell(row=1, column=ci, value=name).font = Font(name=F, bold=True)
    for ri, v in enumerate(vals, start=2):
        ls.cell(row=ri, column=ci, value=v).font = Font(name=F)
ls.sheet_state = "hidden"
list_ref = {}
for ci, (name, vals) in enumerate(LISTS.items(), start=1):
    col = get_column_letter(ci)
    list_ref[name] = f"=Listeler!${col}$2:${col}${len(vals)+1}"

ROWS = 1000


def build(title, cols):
    ws = wb.create_sheet(title)
    ws.freeze_panes = "B2"
    for ci, (key, req, width, note, typ) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=ci, value=key + (" *" if req else ""))
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD_REQ if req else HEAD_OPT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        c.comment = Comment(note, "ALKASU", width=260, height=90)
        col = get_column_letter(ci)
        ws.column_dimensions[col].width = width
        rng = f"{col}2:{col}{ROWS+1}"
        if typ.startswith("list:"):
            dv = DataValidation(type="list", formula1=list_ref[typ[5:]], allow_blank=True,
                                showErrorMessage=typ != "list:birim",
                                errorTitle="Geçersiz değer", error="Listeden bir değer seçin.")
            ws.add_data_validation(dv); dv.add(rng)
        elif typ == "int":
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                showErrorMessage=True, errorTitle="Geçersiz miktar", error="Tam sayı girin (0 veya üzeri).")
            ws.add_data_validation(dv); dv.add(rng)
        elif typ == "money":
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                showErrorMessage=True, errorTitle="Geçersiz tutar", error="0 veya üzeri bir tutar girin.")
            ws.add_data_validation(dv); dv.add(rng)
        fmt = {"money": '#,##0.00', "int": '0', "text": '@'}.get(typ)
        for r in range(2, ROWS + 2):
            cell = ws.cell(row=r, column=ci)
            cell.font = Font(name=F)
            if fmt:
                cell.number_format = fmt
    ws.row_dimensions[1].height = 30
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    return ws


# ---- Açıklama
guide = wb.create_sheet("Açıklama", 0)
guide.column_dimensions["A"].width = 22
guide.column_dimensions["B"].width = 95
lines = [
    ("ALKASU – İçe aktarma şablonu", None),
    ("Sürüm", "1.0 · 29.09.2026"),
    ("", None),
    ("Nasıl doldurulur", None),
    ("1", "Ürünleri 'Ürünler' sayfasına, her ürün bir satır olacak şekilde girin. Koyu mavi başlıklar zorunlu, açık mavi başlıklar isteğe bağlıdır."),
    ("2", "Başlığın üzerine gelince o sütunun açıklaması görünür. Başlık adlarını ve sütun sırasını değiştirmeyin."),
    ("3", "Stok ve fiyatlar TEMEL BİRİM üzerinden girilir. Paket/koli için yalnızca 'içerik' (kaç temel birim) ve fiyat girilir."),
    ("4", "Fiyatlar KDV dahil ve Türk lirası olarak girilir. Ondalık ayırıcı Excel'in ayarına göre virgül veya noktadır."),
    ("5", "Damacana gibi depozitolu ürünler için önce boş kap ürününü ayrı bir satırda tanımlayın (bos_kap_mi = E), sonra dolu ürünün 'bos_kap_urun_kodu' sütununa bu kodu yazın."),
    ("6", "Müşteri geçişi için 'Müşteriler' ve 'Müşteri Kapları' sayfalarını kullanın. Şimdilik yalnızca ürünler yeterlidir."),
    ("7", "Yükleme sırasında ALKASU önce önizleme gösterir. Hatalı satır varsa hiçbir kayıt yapılmaz; hataları düzeltip tekrar yüklersiniz."),
    ("", None),
    ("Örnek", "Doldurulmuş 3 örnek satır için 'Örnek' sayfasına bakın: boş damacana, dolu damacana ve paket/kolili su. Örnek sayfası içe aktarılmaz."),
]
r = 1
for a, b in lines:
    ca = guide.cell(row=r, column=1, value=a)
    ca.font = Font(name=F, bold=(b is None), size=14 if r == 1 else 11)
    if b is not None:
        cb = guide.cell(row=r, column=2, value=b)
        cb.font = Font(name=F)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(vertical="top")
    r += 1


build("Ürünler", PRODUCTS)
build("Müşteriler", CUSTOMERS)
build("Müşteri Kapları", CONTAINERS)

# ---- Örnek sayfası (tam sütunlarla, kurgusal veri)
EX = [
    {"urun_kodu": "BK-001", "urun_adi": "Örnek Marka Boş Damacana", "marka": "Örnek Marka", "kategori": "Boş kap",
     "temel_birim": "Adet", "kdv_orani": 20, "acilis_stogu": 40, "bos_kap_mi": "E", "aktif": "E"},
    {"urun_kodu": "DM-001", "urun_adi": "Örnek Marka 19 L Damacana", "marka": "Örnek Marka", "kategori": "Damacana",
     "temel_birim": "Damacana", "barkod": "8690000000011", "satis_fiyati": 120, "kdv_orani": 1, "alis_fiyati": 85,
     "acilis_stogu": 60, "kritik_stok": 15, "depozitolu": "E", "depozito_tutari": 250, "bos_kap_urun_kodu": "BK-001",
     "bos_kap_mi": "H", "aktif": "E"},
    {"urun_kodu": "SU-005", "urun_adi": "Örnek Marka 0,5 L Su", "marka": "Örnek Marka", "kategori": "Su",
     "temel_birim": "Şişe", "barkod": "8690000000028|8690000000035", "satis_fiyati": 10, "paket_adi": "Paket",
     "paket_icerik": 6, "paket_satis_fiyati": 55, "koli_adi": "Koli", "koli_icerik": 24, "koli_satis_fiyati": 200,
     "kdv_orani": 1, "alis_fiyati": 6.25, "acilis_stogu": 480, "kritik_stok": 96, "depozitolu": "H",
     "bos_kap_mi": "H", "aktif": "E"},
]
exs = wb.create_sheet("Örnek")
exs.sheet_properties.tabColor = "BFBFBF"
for ci, (key, req, width, *_ ) in enumerate(PRODUCTS, start=1):
    h = exs.cell(row=1, column=ci, value=key + (" *" if req else ""))
    h.font = Font(name=F, bold=True, color="FFFFFF"); h.fill = HEAD_REQ if req else HEAD_OPT
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); h.border = BORDER
    exs.column_dimensions[get_column_letter(ci)].width = width
    for ri, ex in enumerate(EX, start=2):
        c = exs.cell(row=ri, column=ci, value=ex.get(key))
        c.font = Font(name=F, italic=True, color="404040"); c.border = BORDER
exs.row_dimensions[1].height = 30
exs.freeze_panes = "B2"
exs.cell(row=6, column=1, value="Bu sayfa yalnızca örnektir ve içe aktarılmaz. Ürün adları ve fiyatlar kurgusaldır.").font = Font(name=F, bold=True, color="C00000")

wb._sheets.remove(ls); wb._sheets.append(ls)
for _ws in wb.worksheets:
    _ws.page_setup.orientation = "landscape"
    _ws.page_setup.fitToWidth = 1
    _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(OUT + "urun_ice_aktarma_sablonu.xlsx")

# CSV (yalnızca ürünler) — UTF-8 BOM, ; ayırıcı (Türkçe Excel)
with open(OUT + "urun_ice_aktarma_sablonu.csv", "w", encoding="utf-8-sig", newline="") as f:
    csv.writer(f, delimiter=";").writerow([k for k, *_ in PRODUCTS])
print("ok")
